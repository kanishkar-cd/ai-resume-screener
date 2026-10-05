"""Regression tests for the deterministic (local) resume parsing pipeline.

Each case mirrors a layout seen in real resumes: two-column sidebars, wrapped
bullets, right-aligned dates, label/value skill tables, issuer-suffixed
certifications and Word list/table templates.
"""

from pathlib import Path

import fitz
from docx import Document

from app.services.extractors.resume_extractor import ResumeExtractor
from app.services.parsers.docx_parser import parse_docx
from app.services.parsers.pdf_layout import _normalize_bullet, reflow_plain_text
from app.services.parsers.pdf_parser import parse_pdf
from app.services.pipeline.extraction_pipeline import reconstruct_layout_text


def _pdf(tmp_path: Path, items: list[tuple[float, float, str, float]]) -> Path:
    """Build a one-page PDF from (x, y, text, fontsize) items."""
    path = tmp_path / "resume.pdf"
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    for x, y, text, size in items:
        page.insert_text((x, y), text, fontsize=size, fontname="helv")
    document.save(path)
    document.close()
    return path


# ---------------------------------------------------------------- PDF layout


def test_two_column_pdf_is_read_column_by_column(tmp_path: Path) -> None:
    left = [f"Left column entry {i} about cloud" + " work" * (i % 3) + "." for i in range(8)]
    right = [f"Right column entry {i} about web" + " apps" * (i % 3) + "." for i in range(8)]
    items = [(40, 80 + 14 * i, text, 9) for i, text in enumerate(left)]
    items += [(320, 80 + 14 * i, text, 9) for i, text in enumerate(right)]
    lines = parse_pdf(_pdf(tmp_path, items)).raw_text.splitlines()

    assert lines[:8] == left
    assert lines[8:] == right


def test_sidebar_column_holding_the_name_is_read_first(tmp_path: Path) -> None:
    items = [(320, 40, "PRIYA SHARMA", 9), (320, 54, "priya@example.com", 9)]
    items += [(320, 80 + 14 * i, f"Right column certification entry {i} Udemy", 9) for i in range(8)]
    items += [(40, 200 + 14 * i, f"Left column education entry number {i}", 9) for i in range(8)]
    text = parse_pdf(_pdf(tmp_path, items)).raw_text

    assert text.splitlines()[0] == "PRIYA SHARMA"
    assert text.index("Right column certification entry 7") < text.index("Left column education entry number 0")


def test_wrapped_bullet_is_rejoined_into_one_line(tmp_path: Path) -> None:
    # Base-14 PDF fonts re-encode non-ASCII glyphs, so the fixtures use ASCII bullets/dashes.
    items = [
        (40, 80, "EXPERIENCE", 11),
        (40, 100, "- Built an ASR pre-processing pipeline with WhisperX, voice activity detection and speaker", 9),
        (48, 112, "diarization for downstream NLP processing.", 9),
        (40, 124, "- Reranked hybrid retrieval results using Reciprocal Rank Fusion.", 9),
    ]
    lines = parse_pdf(_pdf(tmp_path, items)).raw_text.splitlines()

    assert lines[1] == (
        "• Built an ASR pre-processing pipeline with WhisperX, voice activity detection and speaker "
        "diarization for downstream NLP processing."
    )
    assert lines[2].startswith("• Reranked")


def test_title_and_right_aligned_dates_share_one_row(tmp_path: Path) -> None:
    items = [
        (40, 80, "Software Engineer | Cloud Destinations", 10),
        (470, 80, "Jun 2026 - Present", 10),
        (40, 94, "- Built multi-agent accelerators.", 9),
    ]
    lines = parse_pdf(_pdf(tmp_path, items)).raw_text.splitlines()

    assert lines[0] == "Software Engineer | Cloud Destinations  Jun 2026 - Present"
    assert lines[1] == "• Built multi-agent accelerators."


def test_large_name_banner_is_not_merged_into_contact_row(tmp_path: Path) -> None:
    items = [(40, 60, "aswin@example.com | +91 9344081155", 9), (380, 62, "ASWIN SURIYA C", 22)]
    text = parse_pdf(_pdf(tmp_path, items)).raw_text

    assert "ASWIN SURIYA C" in text.splitlines()
    assert ResumeExtractor().extract(text)["candidate_name"] == "Aswin Suriya C"


def test_symbol_font_bullets_are_normalized() -> None:
    assert _normalize_bullet(" Built REST APIs") == ("• Built REST APIs", True)
    assert _normalize_bullet("- Built REST APIs") == ("• Built REST APIs", True)
    assert _normalize_bullet("Built REST APIs") == ("Built REST APIs", False)


def test_ocr_text_reflow_joins_wrapped_lines_only() -> None:
    text = (
        "PROJECT EXPERIENCE\n"
        "PLC based Smart Parking system with IoT\n"
        "Developed a PLC based Smart Parking System with IoT using Ultrasonic Sensor, LED Display, PLC\n"
        "Programming and IoT Platform.\n"
        "Frontend E-commerce Platform\n"
        "Developed a responsive e-commerce frontend website using HTML, CSS and JavaScript today.\n"
    )
    lines = reflow_plain_text(text).splitlines()

    assert lines == [
        "PROJECT EXPERIENCE",
        "PLC based Smart Parking system with IoT",
        "Developed a PLC based Smart Parking System with IoT using Ultrasonic Sensor, LED Display, PLC Programming and IoT Platform.",
        "Frontend E-commerce Platform",
        "Developed a responsive e-commerce frontend website using HTML, CSS and JavaScript today.",
    ]


def test_pipe_separated_lines_are_not_treated_as_columns() -> None:
    text = "\n".join(["Languages | Java | Python", "Tools | Git | Docker", "Phone | Email | GitHub", "Summary line"])
    assert reconstruct_layout_text(text) == text


# ---------------------------------------------------------------- experience


def test_experience_header_fields_are_classified() -> None:
    block = (
        "Software Engineer | Cloud Destinations  June 2026 – Present\n"
        "• Developing multi-agent AI accelerators.\n"
        "Software Engineer Intern — AI/ML Systems  —  KCIRI  Dec 2024 – May 2025\n"
        "• Reduced inference latency by ~30% by parallelising I/O and model execution threads, eliminating\n"
        "Tamil Nadu Newsprint and Paper Limited (TNPL) - Trainee  JULY 2025\n"
        "• Analysed DCS architecture.\n"
    )
    items = ResumeExtractor._experience(block)

    assert [(i["title"], i["company"]) for i in items] == [
        ("Software Engineer", "Cloud Destinations"),
        ("Software Engineer Intern", "KCIRI"),
        ("Trainee", "Tamil Nadu Newsprint and Paper Limited (TNPL)"),
    ]
    assert items[0]["is_current"] is True
    assert items[1]["start_date"] == "Dec 2024" and items[1]["end_date"] == "May 2025"
    assert items[2]["employment_type"] == "Trainee"


def test_wrapped_bullet_fragments_never_become_new_jobs() -> None:
    block = (
        "Data Analyst | Cognizant Technology Solutions India Pvt Ltd  Sep 2023 – May 2024\n"
        "Domain: Digital Marketing Analytics\n"
        "• Owned end-to-end KPI definition and reporting for digital marketing campaigns.\n"
        "• Designed scalable Power BI solutions handling 10M+ rows.\n"
        "Tools: MySQL, Power BI, Python (Pandas)\n"
        "QA Engineer  |  RKN Software Solutions Pvt Ltd  Jan 2019 – Apr 2021\n"
        "• Designed and executed automated UI test suites using Selenium WebDriver.\n"
    )
    items = ResumeExtractor._experience(block)

    assert [(i["title"], i["company"]) for i in items] == [
        ("Data Analyst", "Cognizant Technology Solutions India Pvt Ltd"),
        ("QA Engineer", "RKN Software Solutions Pvt Ltd"),
    ]
    assert len(items[0]["responsibilities"]) == 3


def test_title_and_company_on_separate_header_lines() -> None:
    block = (
        "AI-ML Virtual Internship  Apr 2024 – Jun 2024\n"
        "Google for Developers, AICTE EduSkills  Remote\n"
        "• Completed a 10-week virtual internship focused on Machine Learning.\n"
    )
    [item] = ResumeExtractor._experience(block)

    assert item["title"] == "AI-ML Virtual Internship"
    assert item["company"] == "Google for Developers, AICTE EduSkills"
    assert item["location"] == "Remote"
    assert item["employment_type"] == "Internship"


# ---------------------------------------------------------------- projects


def test_project_headers_with_links_and_dates() -> None:
    block = (
        "Disaster Relief Management System:(DisasterRelief)  [Link]  Feb 2026\n"
        "• Developed a scalable disaster relief coordination platform.\n"
        "Tech Stack: Next.js | Node.js | Express.js | MongoDB\n"
        "SustainTrack.me -  2024\n"
        "Developed a web platform to track and reduce carbon footprint through eco-friendly activities.\n"
        "Devops-Based Ecommerce Website\n"
        "Built a full-stack E-Commerce web app with cart, order processing, and PDF billing features.\n"
    )
    projects = ResumeExtractor._projects(block)

    assert [p["name"] for p in projects] == [
        "Disaster Relief Management System (DisasterRelief)",
        "SustainTrack.me",
        "Devops-Based Ecommerce Website",
    ]
    assert {"Next.js", "Node.js", "Express.js", "MongoDB"} <= set(projects[0]["technologies"])
    assert projects[0]["description"].startswith("Disaster Relief Management System")


def test_bullet_titled_projects_with_wrapped_titles() -> None:
    block = (
        "• AI-Enabled Agro-Climate Intelligence & Adaptive Farm\n"
        "Decision Support System: Built an AI-powered agro-climate intelligence platform for soil moisture prediction.\n"
        "Tech Stack: Python, FastAPI, Scikit-learn, Firebase\n"
        "• CollabDesk – Collaborative Idea & Task Management Platform: Developed a MERN stack web application.\n"
        "Tech Stack: MERN Stack, JWT Authentication, REST APIs\n"
    )
    projects = ResumeExtractor._projects(block)

    assert [p["name"] for p in projects] == [
        "AI-Enabled Agro-Climate Intelligence & Adaptive Farm Decision Support System",
        "CollabDesk – Collaborative Idea & Task Management Platform",
    ]
    assert "Firebase" in projects[0]["technologies"]


# ---------------------------------------------------------------- education


def test_education_institution_and_degree_on_alternating_lines() -> None:
    block = (
        "Kumaraguru College of Technology  Coimbatore, India\n"
        "Bachelor of Engineering in Electronics and Instrumentation; CGPA: 8.83  2022 - 2026\n"
        "Rahmania Matric Higher Secondary School  Kamuthi, India\n"
        "Higher Secondary School Certificate (HSC) - 95.8%  2020 - 2022\n"
        "Rahmania Garden Matric School  Kottaimedu, India\n"
        "Secondary School Leaving Certificate - 94.5%  2019 - 2020\n"
    )
    entries = ResumeExtractor._education(block)

    assert [(e["degree"], e["institution"]) for e in entries] == [
        ("Bachelor of Engineering", "Kumaraguru College of Technology"),
        ("Higher Secondary (12th)", "Rahmania Matric Higher Secondary School"),
        ("Secondary School (10th)", "Rahmania Garden Matric School"),
    ]
    assert entries[0]["year"] == "2022-2026"
    assert entries[0]["field_of_study"] == "Electronics and Instrumentation Engineering"


def test_education_recognizes_hse_and_spelled_out_degrees() -> None:
    block = (
        "Sri Eshwar College of Engineering  B.E(CSE)  CGPA: 7.5  2023-2027\n"
        "Palani Gounder Higher Secondary School  HSE  Percentage: 81.2%  2021-2023\n"
    )
    entries = ResumeExtractor._education(block)

    assert [e["degree"] for e in entries] == ["Bachelor of Engineering", "Higher Secondary (12th)"]


# ---------------------------------------------------------------- certifications


def test_certifications_attach_issuers_and_keep_dashed_names() -> None:
    block = (
        "Learn Java Programming (Master)  UDEMY  2024\n"
        "SQL - Basics  Skill Rack  2024\n"
        "C Programming – Udemy | 2023\n"
        "• IoT Devices (University of Illinois Urbana-Champaign)\n"
        "Issued Mar 2025 - Credential ID: C99Y3YPINL74\n"
        "AWS Certified Cloud Practitioner  ·  MAAS INFO Certified Python Programming — Foundations\n"
    )
    certs = ResumeExtractor._certifications(block)

    assert certs == [
        "Learn Java Programming (Master) (UDEMY)",
        "SQL - Basics (Skill Rack)",
        "C Programming (Udemy)",
        "IoT Devices (University of Illinois Urbana-Champaign)",
        "AWS Certified Cloud Practitioner",
        "MAAS INFO Certified Python Programming — Foundations",
    ]


# ---------------------------------------------------------------- skills & contact


def test_skill_category_labels_without_colons_are_dropped() -> None:
    block = (
        "Languages  Java | HTML | CSS\n"
        "Core Concepts – DBMS\n"
        "Generative AI / Agentic AI  LangGraph · RAG · Prompt Engineering\n"
        "BI & Reporting  Power BI (DAX, KPI Design), Advanced Excel\n"
    )
    skills = ResumeExtractor._skills(block, "")

    assert skills == [
        "Java", "HTML", "CSS", "DBMS", "LangGraph", "RAG", "Prompt Engineering",
        "Power BI", "DAX", "KPI Design", "Advanced Excel",
    ]


def test_phone_repairs() -> None:
    assert ResumeExtractor._extract_phone("Phone: +9 1 8778173009 | Email: a@b.com") == "+91 8778173009"
    assert ResumeExtractor._extract_phone("+91 95973744661 harshini@example.com") == "+91 9597374466"
    assert ResumeExtractor._extract_phone("+91 9344081155") == "+91 9344081155"


def test_location_label_does_not_capture_next_line() -> None:
    text = "HARSHINI CS\nharshini@example.com\nLanguages:\nAddress:\nEnglish,Tamil.\n"
    extracted = ResumeExtractor().extract(text)

    assert extracted["location"] is None
    assert extracted["candidate_name"] == "Harshini CS"


# ---------------------------------------------------------------- DOCX


def test_docx_keeps_document_order_bullets_and_header_name(tmp_path: Path) -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "PRIYA SHARMA"
    document.add_paragraph("EXPERIENCE")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Backend Developer | Zoho Corporation"
    table.cell(0, 1).text = "Jan 2023 – Present"
    document.add_paragraph("Built REST APIs in FastAPI.", style="List Bullet")
    document.add_paragraph("EDUCATION")
    path = tmp_path / "resume.docx"
    document.save(path)

    lines = parse_docx(path).raw_text.splitlines()

    assert lines == [
        "PRIYA SHARMA",
        "EXPERIENCE",
        "Backend Developer | Zoho Corporation | Jan 2023 – Present",
        "• Built REST APIs in FastAPI.",
        "EDUCATION",
    ]
